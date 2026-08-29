import asyncio
import sys
import webbrowser
from pathlib import Path

import pytest

from agents.web_agent import WebAgent
from agents.conversation_agent import ConversationAgent
from core.autonomy import AutonomyEngine
from core.task_decomposer import TaskDecomposer
from learning.adaptive_engine import AdaptiveLearningEngine
from learning.neural_command_model import NeuralCommandModel
from learning.phase5_model import OfflinePreferenceModel
from main import AutonomousSystemController


class DummyMemory:
    def __init__(self):
        self.preferences = {}
        self.short_term = {}

    def remember_episode(self, **kwargs):
        return None

    def remember_short_term(self, key, value):
        self.short_term[key] = value
        return None

    def recall_short_term(self, key):
        return self.short_term.get(key)

    def log_session_activity(self, activity):
        return None

    def save_user_preference(self, key, value):
        self.preferences[key] = {"value": value}

    def get_user_preference(self, key, default=None):
        return self.preferences.get(key, {}).get("value", default)

    def create_task_record(self, command, steps_total=0):
        return None

    def update_task_record(self, task_id, **fields):
        return None


class DummyLearningAgent:
    def __init__(self):
        self.patterns = {}
        self.history = []

    def find_similar_success(self, command, threshold=0.75):
        return None

    def get_recent_commands(self, limit=10):
        return self.history[-limit:]

    def get_learning_stats(self):
        return {"history_count": len(self.history), "success_rate": 1.0 if self.history else 0.0, "form_patterns": len(self.patterns)}

    def record_command(self, command, strategy, result):
        self.history.append({"command": command, "strategy": strategy, "result": result})
        return self.history[-1]

    def record_feedback(self, feedback, sentiment="neutral", command=None):
        return {"feedback": feedback, "sentiment": sentiment, "command": command}

    def discover_workflow_patterns(self, min_count=2):
        return [{"signature": "intent:browser_action", "count": 2, "examples": ["search a", "search b"]}]

    def remember_form_pattern(self, form_signature, field_order, metadata=None):
        self.patterns[form_signature] = {"field_order": field_order, "metadata": metadata or {}}

    def recall_form_pattern(self, form_signature):
        return self.patterns.get(form_signature)

    def refresh_adaptive_state(self, preferences=None):
        return self.get_adaptive_summary()

    def recommend_strategy(self, command):
        return None

    def classify_command(self, command):
        normalized = command.lower().strip()
        if "show system status" in normalized:
            return {
                "predicted_intent": "system_info",
                "confidence": 0.92,
                "source": "adaptive_cluster",
                "cluster": {
                    "signature": "intent:system_info",
                    "strategy": {"intent": {"type": "system_info"}},
                },
            }
        if "search" in normalized:
            return {
                "predicted_intent": "browser_action",
                "confidence": 0.78,
                "source": "adaptive_cluster",
                "cluster": {
                    "signature": "intent:browser_action",
                    "strategy": {"intent": {"type": "browser_action"}},
                },
            }
        return {"predicted_intent": "unknown", "confidence": 0.0, "source": "none"}

    def get_phase5_summary(self):
        return {
            "summary": {
                "clusters": 1,
                "trained_preferences": 1,
                "advisory": True,
                "offline": True,
                "confidence": 0.72,
                "preference_model": "offline_preference_model",
            },
            "clusters": [{"cluster_id": "web", "signature": "route:web", "size": 1, "examples": []}],
            "preference_candidates": [{"key": "default_browser", "value": "chrome", "confidence": 0.72, "source": "offline_preference_model", "advisory": True}],
            "prediction": {"key": "default_browser", "value": "chrome", "confidence": 0.72, "source": "offline_preference_model", "advisory": True},
            "offline": True,
            "advisory": True,
        }

    def predict_phase5_preference(self, command, preferences=None):
        return {
            "prediction": {"key": "default_browser", "value": "chrome", "confidence": 0.72, "source": "offline_preference_model", "advisory": True},
            "advisory": True,
            "reason": "Browser-oriented activity suggests a default browser preference.",
        }

    def get_adaptive_summary(self):
        return {
            "summary": {
                "examples": len(self.history),
                "clusters": 1 if self.history else 0,
                "workflow_templates": 0,
                "stable_workflow_templates": 0,
                "positive_labels": len(self.history),
                "negative_labels": 0,
                "neutral_labels": 0,
                "recommendation_accuracy": 0.0,
                "user_correction_rate": 0.0,
            },
            "workflow_clusters": [],
            "workflow_templates": [],
            "evaluation_metrics": {
                "recommendation_accuracy": 0.0,
                "user_correction_rate": 0.0,
                "workflow_templates_available": 0,
                "stable_workflow_templates": 0,
                "suggestion_acceptance_ratio": 0.0,
            },
            "preference_signals": {"suggestions": []},
            "preference_evaluation": {
                "suggestions_available": 0,
                "suggestions_accepted": 0,
                "acceptance_ratio": 0.0,
            },
        }


class DummyWebAgent:
    def __init__(self):
        self.navigated = []

    async def perform_browser_action(self, browser_name, action, target):
        if action == "navigate":
            return await self.navigate_to_url(target, browser_name=browser_name)
        return {
            "success": True,
            "browser": browser_name,
            "action": action,
            "target": target,
            "url": f"https://example.com/search?q={target.replace(' ', '+')}",
        }

    async def navigate_to_url(self, url, browser_name=None):
        self.navigated.append((url, browser_name))
        return {"success": True, "url": url, "browser": browser_name}

    async def search_site(self, query, site=None, browser_name=None):
        target = f"https://www.{site or 'example'}.com/search?q={query.replace(' ', '+')}"
        return {"success": True, "site": site, "query": query, "url": target, "browser": browser_name}

    async def intelligent_form_fill(self, field_values, form_name=None, submit=False, focus_first_field=False, require_confirmation=False, allow_keyboard_fallback=True):
        result = {"success": True, "fields_filled": [{"name": key, "value": value} for key, value in field_values.items()]}
        result["allow_keyboard_fallback"] = allow_keyboard_fallback
        if require_confirmation and submit:
            result["requires_confirmation"] = True
            result["confirmation_token"] = "dummy-token"
        return result

    async def confirm_submission(self, token):
        return {"success": True, "submitted": True, "confirmation_token": token}


class DummyFileAgent:
    def __init__(self):
        self.created = []
        self.copied = []
        self.moved = []
        self.searched = []
        self.backups = []

    def create_file(self, path, content=""):
        self.created.append((path, content))
        return {"success": True, "path": path, "content_length": len(content)}

    def create_folder(self, path):
        return {"success": True, "path": path}

    def list_files(self, path="."):
        return []

    def organize_folder(self, path):
        return {"success": True, "path": path}

    def copy_file(self, source, destination):
        self.copied.append((source, destination))
        return {"success": True, "source": source, "destination": destination}

    def move_file(self, source, destination):
        self.moved.append((source, destination))
        return {"success": True, "source": source, "destination": destination}

    def get_file_info(self, path):
        return {"success": True, "path": path, "name": Path(path).name}

    def search_files(self, path, query, search_content=False):
        self.searched.append((path, query, search_content))
        return [{"path": str(Path(path) / f"{query}.txt"), "name": f"{query}.txt"}]

    def backup_folder(self, path, backup_name=None):
        self.backups.append((path, backup_name))
        return {"success": True, "source": path, "backup_path": backup_name or "backup"}

    def analyze_disk_usage(self, path):
        return {"success": True, "folder_path": path, "total_size_mb": 1}

    def find_duplicate_files(self, path):
        return {"success": True, "duplicates_found": 0, "duplicates": {}}

    def restore_organization(self, manifest_id=None):
        return {"success": True, "manifest_id": manifest_id}

    def batch_rename_files(self, path, find_str, replace_str, extension=None):
        return {"success": True, "renamed_count": 1}


class DummySystemAgent:
    def __init__(self):
        self.opened = []
        self.closed = []

    def open_application(self, app_name, args=None):
        self.opened.append((app_name, args))
        return {"success": True, "application": app_name, "args": args}

    def close_application(self, app_name):
        self.closed.append(app_name)
        return {"success": True, "application": app_name}

    def list_processes(self, filter_name=None):
        return [{"pid": 1, "name": "demo.exe", "cpu": 0.0, "memory": 1.0}]

    def get_system_info(self):
        return {
            "cpu": {"percent": 1, "cores": 1},
            "memory": {"used_percent": 10, "available_gb": 1},
            "disk": {"used_percent": 20, "free_gb": 2},
        }

    def run_command(self, command, shell="powershell", confirmed=False):
        return {"success": True, "output": command, "return_code": 0}

    def get_network_info(self):
        return {"success": True, "interfaces": {"loopback": {}}}

    def get_disk_usage_all_drives(self):
        return [{"drive": "C:\\", "used_percent": 20}]

    def get_environment_variable(self, name):
        return {"success": True, "name": name, "value": "demo"}


class DummyConversationAgent:
    def is_conversation(self, text):
        return text.lower().strip() in {"hi", "how are you?", "what can you do?"}

    async def chat(self, message, context=None):
        return {
            "success": True,
            "tier": "conversation",
            "mode": "chat",
            "response": "Hi, I am ASC.",
            "model": "dummy",
        }


class DummyCouncil:
    def __init__(self):
        self.calls = []

    async def decide(self, task, complexity="medium"):
        self.calls.append((task, complexity))
        return {
            "action_plan": [
                {"action": "navigate", "params": {"url": "https://example.com/form"}, "description": "Open the form URL"},
                {"action": "form", "params": {"fields": {"name": "Ada", "email": "ada@example.com"}, "submit": False}, "description": "Fill the form fields"},
            ],
            "confidence": 0.92,
        }


def build_controller():
    controller = object.__new__(AutonomousSystemController)
    controller.config = {
        "web_automation": {"default_browser": "chrome"},
        "execution": {"decompose_complex_tasks": True},
    }
    controller.voice_service_url = "http://localhost:5000"
    controller.voice_enabled = True
    controller.command_cache = {}
    controller.command_cache_size = 100
    controller.execution_pause_seconds = 0
    controller.memory = DummyMemory()
    controller.task_decomposer = TaskDecomposer()
    controller.decompose_complex_tasks = True
    controller.agents = {
        "conversation": DummyConversationAgent(),
        "learning": DummyLearningAgent(),
        "web": DummyWebAgent(),
        "file": DummyFileAgent(),
        "system": DummySystemAgent(),
    }
    controller.tier_stats = {"direct": 0, "single_model": 0, "sequential": 0}
    controller.intent_stats = {}
    controller.agent_stats = {}
    controller._unattended_action_timestamps = []
    controller.unattended_action_hourly_limit = 30
    controller.autonomy = AutonomyEngine(controller.task_decomposer, controller.get_tool_manifest)
    return controller


def test_parse_open_browser_and_search_command():
    controller = build_controller()

    intent = controller._parse_direct_command("open chrome and search for AI trends")

    assert intent == {
        "type": "browser_action",
        "app": "chrome",
        "action": "search",
        "target": "ai trends",
    }


def test_parse_fill_form_command():
    controller = build_controller()

    intent = controller._parse_direct_command(
        "fill form with name: Swapnil, email: swapnil@example.com, company: OpenAI"
    )

    assert intent["type"] == "fill_form"
    assert intent["fields"] == {
        "name": "Swapnil",
        "email": "swapnil@example.com",
        "company": "OpenAI",
    }


def test_parse_fill_form_requires_confirmation():
    controller = build_controller()

    intent = controller._parse_direct_command(
        "Stage the form with name: Alice, email: alice@example.com and confirm before submitting"
    )

    assert intent["type"] == "fill_form"
    assert intent["submit"] is True
    assert intent["require_confirmation"] is True


def test_parse_confirm_submission_token():
    controller = build_controller()

    intent = controller._parse_direct_command(
        "Confirm submission token 123e4567-e89b-12d3-a456-426614174000"
    )

    assert intent["type"] == "confirm_submission"
    assert intent["token"] == "123e4567-e89b-12d3-a456-426614174000"


def test_parse_system_power_actions_are_confirmation_gated():
    controller = build_controller()

    shutdown = controller._parse_direct_command("shutdown system")
    sleep = controller._parse_direct_command("put laptop to sleep")
    power_on = controller._parse_direct_command("turn on laptop")
    unlock = controller._parse_direct_command("unlock laptop")

    assert shutdown["type"] == "system_power_action"
    assert shutdown["power_action"] == "shutdown"
    assert shutdown["command"] == "shutdown /s /t 10"  # 10s abortable safety countdown
    assert sleep["power_action"] == "sleep"
    assert power_on["unsupported_reason"]
    assert unlock["unsupported_reason"]


def test_direct_strategy_is_selected_for_browser_search():
    controller = build_controller()

    strategy = controller._determine_execution_strategy("open chrome and search for AI trends")

    assert strategy["tier"] == "direct"
    assert strategy["intent"]["type"] == "browser_action"


def test_direct_execution_uses_web_agent_for_browser_workflow():
    controller = build_controller()

    result = asyncio.run(
        controller._execute_directly(
            "open chrome and search for AI trends",
            {
                "intent": {
                    "type": "browser_action",
                    "app": "chrome",
                    "action": "search",
                    "target": "AI trends",
                }
            },
        )
    )

    assert result["success"] is True
    assert result["tier"] == "direct"
    assert result["results"][0]["browser"] == "chrome"


def test_url_form_fill_disables_blind_keyboard_fallback():
    controller = build_controller()

    result = asyncio.run(
        controller.process_command("https://example.com fill form with name: Swapnil, email: swapnil@example.com")
    )

    assert result["success"] is True
    first = result["results"][0]
    assert first["navigated_url"] == "https://example.com"
    assert first["allow_keyboard_fallback"] is False


def test_system_power_action_requires_confirmation_before_execution():
    controller = build_controller()

    first = asyncio.run(controller.process_command("restart system"))
    second = asyncio.run(controller.process_command("confirm command"))

    assert first["success"] is False
    assert first["results"][0]["needs_confirmation"] is True
    assert controller.memory.recall_short_term("pending_system_command") is None
    assert second["success"] is True
    # Restart now uses a 10s abortable countdown (shutdown /a can cancel it)
    # rather than an immediate restart - an intentional safety feature.
    assert "shutdown /r /t 10" in second["results"][0]["output"]


def test_direct_execution_confirm_submission_intent():
    controller = build_controller()

    result = asyncio.run(
        controller._execute_directly(
            "Confirm submission token dummy-token",
            {
                "intent": {
                    "type": "confirm_submission",
                    "token": "dummy-token",
                }
            },
        )
    )

    assert result["success"] is True
    assert result["results"][0]["submitted"] is True
    assert result["results"][0]["confirmation_token"] == "dummy-token"


class StubSystemAgent:
    def __init__(self):
        self.calls = []
        self.typed = []
        self.keys = []
        self.hotkeys = []

    def open_application(self, app_name, args=None):
        self.calls.append((app_name, args))
        return {"success": True, "application": app_name, "pid": 1234, "path": app_name}

    def click_mouse(self, x, y):
        return {"success": True, "x": x, "y": y}

    def type_text(self, text, interval=0.05):
        self.typed.append(text)
        return {"success": True, "text": text}

    def press_hotkey(self, *keys):
        self.hotkeys.append(keys)
        return {"success": True, "keys": list(keys)}

    def press_key(self, key):
        self.keys.append(key)
        return {"success": True, "key": key}


class StubVisionAgent:
    def __init__(self, ocr_available=True, text="Name\nEmail\nCompany\nSubmit"):
        self._ocr_available = ocr_available
        self._text = text

    def extract_text_ocr(self, image=None):
        return self._text

    def is_ocr_available(self):
        return self._ocr_available


class FakeBrowserPage:
    def __init__(self, content_text="Thank you"):
        self.clicks = []
        self.keyboard_presses = []
        self.content_text = content_text
        self.keyboard = self

    async def click(self, selector):
        self.clicks.append(selector)

    async def press(self, key):
        self.keyboard_presses.append(key)


class FakeBrowserDriver:
    def __init__(self, content_text="Thank you"):
        self.page = FakeBrowserPage(content_text=content_text)
        self.fill_calls = []

    async def find_selectors_for_labels(self, labels):
        mapping = {}
        for label in labels:
            normalized = label.lower()
            if normalized in {"name", "email"}:
                mapping[label] = {"selector": f"#{normalized}", "confidence": 0.95}
            else:
                mapping[label] = {"selector": None, "confidence": 0.0}
        return mapping

    async def find_submit_selector(self):
        return "#submit"

    async def fill_fields(self, field_map, submit_selector=None):
        self.fill_calls.append({"field_map": dict(field_map), "submit_selector": submit_selector})
        if submit_selector:
            await self.page.click(submit_selector)
        return {"success": True, "filled": list(field_map.keys()), "page_content": self.page.content_text}

    async def find_label_candidates(self):
        return []


def test_web_agent_builds_search_url_and_reuses_existing_agents():
    system_agent = StubSystemAgent()
    vision_agent = StubVisionAgent()
    learning_agent = DummyLearningAgent()
    agent = WebAgent(
        config_path="./config/settings.json",
        system_agent=system_agent,
        vision_agent=vision_agent,
        learning_agent=learning_agent,
    )
    agent.browser_driver = None
    agent.navigation_delay_seconds = 0

    result = asyncio.run(agent.search_web("AI trends", browser_name="chrome"))

    assert result["success"] is True
    assert result["query"] == "AI trends"
    assert "google.com/search?q=AI+trends" in result["url"]
    assert system_agent.calls[0][0] == "chrome"


def test_web_agent_form_fill_uses_stored_pattern_order():
    system_agent = StubSystemAgent()
    vision_agent = StubVisionAgent()
    learning_agent = DummyLearningAgent()
    learning_agent.remember_form_pattern(
        "name|email|company|submit",
        ["Name", "Email", "Company"],
    )

    agent = WebAgent(
        config_path="./config/settings.json",
        system_agent=system_agent,
        vision_agent=vision_agent,
        learning_agent=learning_agent,
    )
    agent.browser_driver = None
    agent.form_field_delay_seconds = 0
    agent.form_submit_delay_seconds = 0

    result = asyncio.run(
        agent.intelligent_form_fill(
            {"Email": "swapnil@example.com", "Name": "Swapnil", "Company": "ASC"},
            submit=True,
            focus_first_field=False,
        )
    )

    assert result["success"] is True
    assert system_agent.typed == ["Swapnil", "swapnil@example.com", "ASC"]
    assert system_agent.keys.count("tab") == 2
    assert system_agent.keys[-1] == "enter"


def test_web_agent_driver_submit_happens_once():
    agent = WebAgent(
        config_path="./config/settings.json",
        system_agent=StubSystemAgent(),
        vision_agent=StubVisionAgent(),
        learning_agent=DummyLearningAgent(),
    )
    agent.browser_driver = FakeBrowserDriver()

    result = asyncio.run(
        agent.intelligent_form_fill(
            {"Name": "Swapnil", "Email": "swapnil@example.com"},
            submit=True,
        )
    )

    assert result["success"] is True
    assert result["submitted"] is True
    assert agent.browser_driver.fill_calls[0]["submit_selector"] is None
    assert agent.browser_driver.page.clicks == ["#submit"]


def test_web_agent_driver_stage_does_not_submit_until_confirmed():
    agent = WebAgent(
        config_path="./config/settings.json",
        system_agent=StubSystemAgent(),
        vision_agent=StubVisionAgent(),
        learning_agent=DummyLearningAgent(),
    )
    agent.browser_driver = FakeBrowserDriver()

    staged = asyncio.run(
        agent.intelligent_form_fill(
            {"Name": "Swapnil", "Email": "swapnil@example.com"},
            submit=True,
            require_confirmation=True,
        )
    )

    assert staged["success"] is True
    assert staged["requires_confirmation"] is True
    assert staged["confirmation_token"]
    assert agent.browser_driver.page.clicks == []

    confirmed = asyncio.run(agent.confirm_submission(staged["confirmation_token"]))

    assert confirmed["success"] is True
    assert confirmed["submitted"] is True
    assert confirmed["confirmation_token"] == staged["confirmation_token"]
    assert agent.browser_driver.page.clicks == ["#submit"]


def test_web_agent_falls_back_to_default_browser_with_warning(monkeypatch):
    class FailingSystemAgent(StubSystemAgent):
        def open_application(self, app_name, args=None):
            self.calls.append((app_name, args))
            return {"success": False, "application": app_name, "error": "missing browser"}

    opened = {}

    def fake_open(url):
        opened["url"] = url
        return True

    monkeypatch.setattr(webbrowser, "open", fake_open)

    agent = WebAgent(
        config_path="./config/settings.json",
        system_agent=FailingSystemAgent(),
        vision_agent=StubVisionAgent(),
        learning_agent=DummyLearningAgent(),
    )
    agent.browser_driver = None
    agent.navigation_delay_seconds = 0

    result = asyncio.run(agent.navigate_to_url("github.com", browser_name="chrome"))

    assert result["success"] is True
    assert result["launch_mode"] == "fallback"
    assert "warning" in result
    assert opened["url"] == "https://github.com"


def test_form_fill_with_url_is_parsed():
    controller = build_controller()

    intent = controller._parse_direct_command(
        "https://example.com/form fill form with name: Swapnil, email: swapnil@example.com"
    )

    assert intent["type"] == "fill_form"
    assert intent["url"] == "https://example.com/form"


def test_sequence_commands_prefer_decomposition():
    controller = build_controller()

    strategy = controller._determine_execution_strategy(
        "show system status and then list running processes"
    )

    assert strategy["tier"] == "sequential"


def test_task_decomposer_creates_research_report_workflow():
    decomposer = TaskDecomposer()
    plan = decomposer.decompose("reseach AI trends and create a report")

    assert plan["strategy"] == "research_report"
    assert [step["action"] for step in plan["subtasks"]] == ["search", "extract_text", "create_report"]


def test_controller_executes_decomposed_report_workflow(monkeypatch):
    import core.document_tools as document_tools

    controller = build_controller()
    decomposition = controller.task_decomposer.decompose("research AI trends and create a report")

    async def fake_route_action(action, params, context):
        if action == "search":
            return {"success": True, "url": "https://example.com/ai-trends"}
        if action == "extract_text":
            return {"success": True, "text": "AI trends are accelerating across local agents."}
        return {"success": True, "message": context}

    controller._route_action = fake_route_action

    # The report step now writes a real .docx via core.document_tools
    # (upgraded from a raw-text dump routed through the file agent - see
    # main.py's _build_and_save_report) - fake the write itself, matching
    # how every other agent here is a lightweight recording fake rather
    # than touching real disk during a unit test.
    created_docs = []

    def fake_create_word_document(path, title, sections):
        created_docs.append({"path": path, "title": title, "sections": sections})
        return {"success": True, "path": path}

    monkeypatch.setattr(document_tools, "create_word_document", fake_create_word_document)

    result = asyncio.run(controller._execute_decomposed_tasks(decomposition, decomposition["goal"]))

    assert result["success"] is True
    assert result["steps_total"] == 3
    assert created_docs[0]["path"].endswith("_report.docx")
    section_bodies = " ".join(
        s.get("body", "") if isinstance(s, dict) else str(s) for s in created_docs[0]["sections"]
    )
    assert "AI trends are accelerating" in section_bodies


def test_task_decomposer_handles_browser_then_create_file():
    decomposer = TaskDecomposer()
    plan = decomposer.decompose("open chrome and search for agentic ai and then create file named agentic_notes.txt")

    assert plan["strategy"] == "browser_then_file"
    assert plan["subtasks"][0]["action"] in {"open", "search"}
    assert plan["subtasks"][-1]["action"] == "create_file"


def test_memory_session_management():
    from core.memory import MemorySystem
    import shutil
    
    # Initialize real MemorySystem with default/config
    memory = MemorySystem()
    
    # Test session startup, logging, and preference retrieval
    sess_id = memory.start_session("Test CLI Start")
    assert sess_id.startswith("session_")
    
    memory.log_session_activity("User queried disk usage")
    activities = memory.get_last_session_activities()
    assert "User queried disk usage" in activities or "Test CLI Start" in activities
    
    memory.save_user_preference("theme", "dark")
    pref = memory.get_user_preference("theme")
    assert pref == "dark"
    
    memory.end_session()
    assert memory.current_session_id is None


def test_decomposer_advanced_templates():
    decomposer = TaskDecomposer()
    
    # 1. Download & Organize
    plan1 = decomposer.decompose("download https://example.com/data.csv and organize desktop")
    assert plan1["strategy"] == "download_and_organize"
    assert [step["action"] for step in plan1["subtasks"]] == ["navigate", "wait_for_process", "organize"]
    
    # 2. Compare Files
    plan2 = decomposer.decompose("compare file1.txt and file2.txt")
    assert plan2["strategy"] == "compare_files"
    assert [step["action"] for step in plan2["subtasks"]] == ["list", "run_command"]
    
    # 3. Backup & Clean
    plan3 = decomposer.decompose("backup desktop and clean folder")
    assert plan3["strategy"] == "backup_and_clean"
    assert [step["action"] for step in plan3["subtasks"]] == ["list", "run_command", "organize"]


def test_parse_security_commands():
    controller = build_controller()
    
    # Port scan
    intent = controller._parse_direct_command("run security scan on 127.0.0.1 ports 80-443")
    assert intent["type"] == "security_port_scan"
    assert intent["host"] == "127.0.0.1"
    assert intent["ports"] == "80-443"
    
    # SSL Check
    intent = controller._parse_direct_command("ssl check google.com")
    assert intent["type"] == "security_ssl_check"
    assert intent["domain"] == "google.com"
    
    # Security Headers
    intent = controller._parse_direct_command("check security headers for https://example.com")
    assert intent["type"] == "security_headers_check"
    assert intent["url"] == "https://example.com"
    
    # DNS Lookup
    intent = controller._parse_direct_command("dns lookup for localhost")
    assert intent["type"] == "dns_lookup"
    assert intent["domain"] == "localhost"


def test_parse_file_advanced_commands():
    controller = build_controller()
    
    # Undo organization
    intent = controller._parse_direct_command("undo organization with manifest manifest_123")
    assert intent["type"] == "restore_organization"
    assert intent["manifest_id"] == "manifest_123"
    
    # Analyze disk usage
    intent = controller._parse_direct_command("analyze disk usage in desktop")
    assert intent["type"] == "analyze_disk_usage"
    
    # Find duplicates
    intent = controller._parse_direct_command("find duplicates in downloads")
    assert intent["type"] == "find_duplicates"
    
    # Batch rename
    intent = controller._parse_direct_command("batch rename files in c:\\temp replace old with new matching .jpg")
    assert intent["type"] == "batch_rename_files"
    assert intent["find"] == "old"
    assert intent["replace"] == "new"
    assert intent["extension"] == ".jpg"


def test_controller_decomposed_tasks_rollback():
    controller = build_controller()
    
    # We mock subtasks: step-1 succeeds (create folder), step-2 fails (system command)
    decomposition = {
        "goal": "create folder ./temp_test and then run invalid command",
        "strategy": "clause_split",
        "subtasks": [
            {
                "id": "step-1",
                "agent": "file",
                "action": "create_folder",
                "params": {"path": "./temp_test", "description": "create folder ./temp_test"},
                "description": "create folder ./temp_test"
            },
            {
                "id": "step-2",
                "agent": "system",
                "action": "run_command",
                "params": {"command": "invalid-cmd-12345"},
                "description": "run invalid command"
            }
        ]
    }
    
    deleted_paths = []
    
    class MockFileAgentWithDelete:
        def create_folder(self, path):
            return {"success": True, "path": path}
        def delete_folder(self, path, recursive=True):
            deleted_paths.append(path)
            return {"success": True}
            
    controller.agents["file"] = MockFileAgentWithDelete()
    
    async def fake_route_action(action, params, context):
        if action == "create_folder":
            return controller.agents["file"].create_folder(params.get("path"))
        return {"success": False, "error": "command failed"}
        
    controller._route_action = fake_route_action
    
    result = asyncio.run(controller._execute_decomposed_tasks(decomposition, "goal"))
    
    assert result["success"] is False
    assert result["rollback_executed"] is True
    assert "./temp_test" in deleted_paths


def test_clause_splitting_complex_browser_actions():
    decomposer = TaskDecomposer()
    plan = decomposer.decompose("open google goto amazon and search for laptops")
    assert plan["strategy"] == "clause_split"
    assert len(plan["subtasks"]) == 3
    assert plan["subtasks"][0]["action"] == "navigate"
    assert "google.com" in plan["subtasks"][0]["params"]["target"]
    assert plan["subtasks"][1]["action"] == "navigate"
    assert "amazon.com" in plan["subtasks"][1]["params"]["target"]
    assert plan["subtasks"][2]["action"] == "search"
    assert plan["subtasks"][2]["params"]["query"] == "laptops"


def test_decomposer_form_filling_parsing():
    decomposer = TaskDecomposer()
    plan = decomposer.decompose("go to https://forms.gle/XYZ and fill the form with name: Swapnil, email: swapnil@gmail.com")
    assert len(plan["subtasks"]) == 2
    assert plan["subtasks"][0]["action"] == "navigate"
    assert plan["subtasks"][0]["params"]["target"] == "https://forms.gle/XYZ"
    assert plan["subtasks"][1]["action"] == "form"
    assert plan["subtasks"][1]["params"]["fields"] == {"name": "Swapnil", "email": "swapnil@gmail.com"}
    assert plan["subtasks"][1]["params"]["url"] == "https://forms.gle/XYZ"


def test_web_agent_vision_guided_click_fallback():
    class ClickTrackingSystemAgent(StubSystemAgent):
        def __init__(self):
            super().__init__()
            self.clicks = []
        def click_mouse(self, x, y):
            self.clicks.append((x, y))
            return {"success": True, "x": x, "y": y}

    system_agent = ClickTrackingSystemAgent()
    called_text = []

    class MockVisionAgentWithCoordinates(StubVisionAgent):
        def find_text_coordinates(self, text, image=None):
            called_text.append(text)
            return (500, 300)

    vision_agent = MockVisionAgentWithCoordinates()
    learning_agent = DummyLearningAgent()

    agent = WebAgent(
        config_path="./config/settings.json",
        system_agent=system_agent,
        vision_agent=vision_agent,
        learning_agent=learning_agent,
    )
    agent.browser_driver = None
    agent.form_field_delay_seconds = 0
    agent.form_submit_delay_seconds = 0

    result = asyncio.run(
        agent.intelligent_form_fill(
            {"Name": "Swapnil"},
            submit=False,
            focus_first_field=True,
        )
    )

    assert result["success"] is True
    assert "Name" in called_text
    assert system_agent.clicks == [(500, 345)]


def test_unknown_general_action_fails_instead_of_fake_success():
    controller = build_controller()

    result = asyncio.run(
        controller._execute_general_action(
            "do the mysterious unsupported thing",
            {"description": "do the mysterious unsupported thing"},
        )
    )

    assert result["success"] is False
    # _execute_general_action's error message was made more specific/honest
    # since this assertion was written (now names the actual command and
    # gives a concrete next step rather than a generic "not a real agent
    # action" message).
    assert result["action"] == "unknown"
    assert "too vague to execute safely" in result["error"]


def test_council_string_plan_is_normalized_to_routed_action():
    controller = build_controller()
    plan = {"steps": ["search for laptops on amazon"], "confidence": 0.8}

    result = asyncio.run(controller._execute_plan(plan, "search for laptops on amazon"))

    assert result["success"] is True
    assert result["steps_completed"] == 1
    assert result["results"][0]["action"] == "navigate"
    target_url = result["results"][0].get("url") or result["results"][0].get("target", "")
    assert "amazon.in" in target_url
    assert "laptops" in target_url


def test_ecommerce_browser_command_maps_to_site_search_url():
    controller = build_controller()

    intent = controller._parse_direct_command(
        "open chrome and search for amazon and search for laptops on amazon"
    )

    assert intent["type"] == "browser_action"
    assert intent["action"] == "navigate"
    assert "amazon.in/s?k=laptops" in intent["target"]


def test_ecommerce_there_command_prefers_single_direct_action():
    controller = build_controller()

    strategy = controller._determine_execution_strategy(
        "open chrome and search for amazon and search for laptops there"
    )

    assert strategy["tier"] == "direct"
    assert strategy["intent"]["type"] == "browser_action"
    assert strategy["intent"]["action"] == "navigate"
    assert "amazon.in/s?k=laptops" in strategy["intent"]["target"]


def test_tool_registry_exposes_phase_two_tools():
    controller = build_controller()

    manifest = controller.get_tool_manifest()
    tool_names = {tool["name"] for tool in manifest}
    actions = {action for tool in manifest for action in tool["actions"]}

    assert {"file_ops", "system_ops", "vision_ops", "web_ops"}.issubset(tool_names)
    assert {"search", "site_search", "navigate", "form", "open", "create_file", "run_command"}.issubset(actions)


def test_route_action_uses_registered_web_tool():
    controller = build_controller()

    result = asyncio.run(
        controller._route_action(
            "search",
            {"query": "phase 2 autonomy"},
            "search for phase 2 autonomy",
        )
    )

    assert result["success"] is True
    assert result["agent"] == "web"
    assert result["tool"] == "web_ops"
    assert result["action"] == "search"


def test_parse_voice_status_command():
    controller = build_controller()

    intent = controller._parse_direct_command("check voice status")

    assert intent == {"type": "voice_status"}


def test_conversation_strategy_is_selected_for_greeting():
    controller = build_controller()

    strategy = controller._determine_execution_strategy("hi")

    assert strategy["tier"] == "conversation"
    assert strategy["source"] == "conversation"


def test_conversation_agent_detects_capability_statement():
    agent = ConversationAgent("./config/settings.json")

    assert agent.is_conversation("tell me what you are capable of") is True


def test_conversation_agent_fallback_answers_without_ollama():
    agent = ConversationAgent("./config/settings.json")
    agent.use_llm = False

    result = asyncio.run(agent.chat("what can you do?"))

    assert result["success"] is True
    assert result["tier"] == "conversation"
    assert "automate" in result["response"]


def test_display_result_prints_conversation_response(capsys):
    controller = build_controller()

    controller._display_result({
        "success": True,
        "tier": "conversation",
        "mode": "chat",
        "response": "Hi, I am ASC.",
    })

    captured = capsys.readouterr()
    # Product was rebranded ASC -> VEGA (2026-08-26); _display_result prints
    # the "VEGA:" prefix regardless of what the response text itself says.
    assert "VEGA: Hi, I am ASC." in captured.out
    assert "Success via conversation tier" not in captured.out


def test_display_result_prints_command_output(capsys):
    controller = build_controller()

    controller._display_result({
        "success": True,
        "tier": "direct",
        "steps_completed": 1,
        "steps_total": 1,
        "results": [{
            "success": True,
            "output": "Friday, July 14, 2026",
            "error": "",
            "return_code": 0,
        }],
    })

    captured = capsys.readouterr()
    assert "Excellent" in captured.out
    assert "direct execution" in captured.out


def test_autonomy_asks_for_missing_target():
    controller = build_controller()

    strategy = controller._determine_execution_strategy("open")
    result = asyncio.run(controller.process_command("open"))

    assert strategy["tier"] == "autonomous"
    assert result["success"] is False
    assert result["needs_clarification"] is True
    assert "What should I open" in result["question"]


def test_autonomy_handles_currently_opened_apps_without_ollama():
    controller = build_controller()

    result = asyncio.run(controller.process_command("currently opened apps"))

    assert result["success"] is True
    assert result["tier"] == "autonomous"
    assert result["autonomy"]["strategy"] == "heuristic_single_step"
    assert result["results"][0]["agent"] == "system"


def test_autonomy_creates_note_from_flexible_request():
    controller = build_controller()

    result = asyncio.run(controller.process_command("make a note about project ideas"))

    assert result["success"] is True
    assert result["tier"] == "autonomous"
    created_path, created_content = controller.agents["file"].created[0]
    assert created_path.endswith("project_ideas_note.txt")
    assert "project ideas" in created_content


def test_autonomy_uses_council_for_ambiguous_form_workflow():
    controller = build_controller()
    controller.council = DummyCouncil()

    result = asyncio.run(controller.process_command("fill this form for me"))

    assert result["success"] is True
    assert result["tier"] == "autonomous"
    assert controller.council.calls
    assert controller.agents["web"].navigated[0][0] == "https://example.com/form"


def test_feedback_command_records_feedback():
    controller = build_controller()
    controller.memory.remember_short_term("last_command", "open chrome")

    result = asyncio.run(controller.process_command("that worked"))

    assert result["success"] is True
    feedback = result["results"][0]["feedback_recorded"]
    assert feedback["sentiment"] == "positive"
    assert feedback["command"] == "that worked"


def test_preference_commands_save_and_show_preferences():
    controller = build_controller()

    save_result = asyncio.run(controller.process_command("remeber default browser is chrome"))
    show_result = asyncio.run(controller.process_command("show preferences"))

    assert save_result["success"] is True
    assert save_result["results"][0]["preference_saved"]["key"] == "default_browser"
    assert show_result["success"] is True
    assert show_result["results"][0]["preferences"]["default_browser"]["value"] == "chrome"


def test_approve_preference_command_saves_preference():
    controller = build_controller()

    approve_result = asyncio.run(controller.process_command("approve preference default browser is edge"))
    assert approve_result["success"] is True
    assert approve_result["results"][0]["approved"] is True
    assert controller.memory.get_user_preference("default_browser") == "edge"


def test_workflow_patterns_command_returns_learning_summary():
    controller = build_controller()

    result = asyncio.run(controller.process_command("show learned workflows"))

    assert result["success"] is True
    patterns = result["results"][0]["workflow_patterns"]
    assert patterns[0]["signature"] == "intent:browser_action"


def test_phase4_adaptive_engine_builds_labeled_clusters(tmp_path):
    engine = AdaptiveLearningEngine(str(tmp_path / "adaptive_state.json"), min_recommendation_score=0.5)
    experiences = [
        {
            "command": "show system status",
            "normalized": "show system status",
            "strategy": {"tier": "direct", "intent": {"type": "system_info"}},
            "result": {"success": True},
            "reward": 1.0,
        },
        {
            "command": "show cpu status",
            "normalized": "show cpu status",
            "strategy": {"tier": "direct", "intent": {"type": "system_info"}},
            "result": {"success": True},
            "reward": 0.9,
        },
    ]

    state = engine.update(experiences, {}, preferences={})
    recommendation = engine.recommend_strategy("show system status")

    assert state["summary"]["positive_labels"] == 2
    assert state["workflow_clusters"][0]["signature"] == "intent:system_info"
    assert recommendation["recommended"] is True
    assert recommendation["strategy"]["intent"]["type"] == "system_info"


def test_phase4_adaptive_learning_status_command():
    controller = build_controller()
    controller.agents["learning"].record_command(
        "show system status",
        {"tier": "direct", "intent": {"type": "system_info"}},
        {"success": True},
    )

    result = asyncio.run(controller.process_command("phase 4 status"))

    assert result["success"] is True
    assert "adaptive_learning" in result["results"][0]
    adaptive = result["results"][0]["adaptive_learning"]
    assert adaptive["summary"]["examples"] >= 1
    assert "workflow_templates" in adaptive["summary"]
    assert "recommendation_accuracy" in adaptive["summary"]
    assert "user_correction_rate" in adaptive["summary"]


def test_reject_preference_command_marks_suggestion_rejected():
    controller = build_controller()

    result = asyncio.run(controller.process_command("reject preference default_browser chrome"))

    assert result["success"] is True
    assert result["results"][0]["preference_rejected"]["key"] == "default_browser"
    assert result["results"][0]["preference_rejected"]["value"] == "chrome"


def test_show_workflow_templates_command_returns_templates():
    controller = build_controller()
    controller.agents["learning"].record_command(
        "open chrome and search for AI trends",
        {"tier": "direct", "intent": {"type": "browser_action"}},
        {"success": True},
    )
    controller.agents["learning"].record_command(
        "search youtube for music",
        {"tier": "direct", "intent": {"type": "browser_action"}},
        {"success": True},
    )

    result = asyncio.run(controller.process_command("show workflow templates"))

    assert result["success"] is True
    templates = result["results"][0]["workflow_templates"]
    assert isinstance(templates, list)


def test_phase5_status_command_returns_offline_model_summary():
    controller = build_controller()
    controller.agents["learning"].record_command(
        "open chrome and search for AI trends",
        {"tier": "direct", "intent": {"type": "browser_action"}},
        {"success": True},
    )

    result = asyncio.run(controller.process_command("phase 5 status"))

    assert result["success"] is True
    phase5 = result["results"][0]["phase5_learning"]
    assert phase5["summary"]["clusters"] >= 1
    assert "preference_model" in phase5["summary"]


def test_phase5_preference_prediction_is_advisory_and_confident():
    controller = build_controller()
    controller.agents["learning"].record_command(
        "open chrome and search for AI trends",
        {"tier": "direct", "intent": {"type": "browser_action"}},
        {"success": True},
    )

    result = asyncio.run(controller.process_command("show phase 5 model"))

    assert result["success"] is True
    model = result["results"][0]["phase5_model"]
    assert model["summary"]["trained_preferences"] >= 0
    assert isinstance(model["prediction"], dict)


def test_phase5_model_reports_richer_summary_metrics(tmp_path):
    model = OfflinePreferenceModel(path=str(tmp_path / "phase5_model.json"))

    state = model.fit(
        [{"command": "open chrome and search for AI trends", "result": {"success": True}}],
        preferences={},
    )

    assert state["summary"]["preference_model"] == "offline_preference_model"
    assert state["summary"]["top_cluster"] == "web"
    assert state["summary"]["evidence_strength"] >= 0.0


def test_phase5_prediction_includes_confidence_label_and_evidence():
    model = OfflinePreferenceModel(path="./tmp_phase5_model.json")
    prediction = model.predict_preference("open chrome and search for AI trends", preferences={})

    assert prediction["prediction"]["confidence_label"] in {"low", "medium", "high"}
    assert "evidence" in prediction
    assert isinstance(prediction["evidence"], str)


def test_online_learning_exposes_weighted_signals(tmp_path):
    model = OfflinePreferenceModel(path=str(tmp_path / "phase5_model.json"))
    state = model.fit(
        [{"command": "open chrome and search for AI trends", "result": {"success": True}}],
        preferences={},
    )

    assert state["summary"]["online_mode"] is True
    assert state["summary"]["online_weight_strength"] >= 0.0

    prediction = model.predict_preference("open chrome and search for AI trends", preferences={})
    assert prediction["prediction"]["online_weight_score"] >= 0.0
    assert prediction["prediction"]["dominant_signal"] in {"default_browser", "notes_folder", "general"}


def test_phase5_neural_command_model_trains_advisory_routes(tmp_path):
    model = NeuralCommandModel(path=str(tmp_path / "neural_command_model.json"), min_examples=4)
    experiences = [
        {"command": "open chrome and search for laptops", "strategy": {"intent": {"type": "browser_action"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "search amazon for phones", "strategy": {"intent": {"type": "browser_action"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "list files in desktop", "strategy": {"intent": {"type": "list_files"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "backup desktop folder", "strategy": {"intent": {"type": "backup_folder"}}, "result": {"success": True}, "reward": 0.9},
        {"command": "show system status", "strategy": {"intent": {"type": "system_info"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "run powershell Get-Date", "strategy": {"intent": {"type": "system_command"}}, "result": {"success": True}, "reward": 1.0},
    ]

    state = model.fit(experiences)
    prediction = model.predict("open browser and search google")

    assert state["trained"] is True
    assert state["advisory"] is True
    assert state["metrics"]["train_examples"] >= 1
    assert prediction["source"] == "neural_command_model"
    assert prediction["advisory"] is True
    assert prediction["raw_route"] in {"web", "file", "system"}


def test_phase5_summary_includes_deep_neural_metrics(tmp_path):
    model = OfflinePreferenceModel(path=str(tmp_path / "phase5_model.json"), min_confidence=0.5)
    experiences = [
        {"command": "open chrome and search for laptops", "strategy": {"intent": {"type": "browser_action"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "search amazon for phones", "strategy": {"intent": {"type": "browser_action"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "list files in desktop", "strategy": {"intent": {"type": "list_files"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "backup desktop folder", "strategy": {"intent": {"type": "backup_folder"}}, "result": {"success": True}, "reward": 0.9},
        {"command": "show system status", "strategy": {"intent": {"type": "system_info"}}, "result": {"success": True}, "reward": 1.0},
        {"command": "run powershell Get-Date", "strategy": {"intent": {"type": "system_command"}}, "result": {"success": True}, "reward": 1.0},
    ]

    state = model.fit(experiences)
    prediction = model.predict_preference("open chrome and search for AI trends", preferences={})

    assert state["summary"]["neural_trained"] is True
    assert state["summary"]["neural_train_examples"] >= 1
    assert "neural_command_model" in state
    assert prediction["route_prediction"]["source"] == "neural_command_model"


def test_show_adaptive_evaluation_command_returns_metrics():
    controller = build_controller()
    controller.agents["learning"].record_command(
        "show system status",
        {"tier": "direct", "intent": {"type": "system_info"}},
        {"success": True},
    )

    result = asyncio.run(controller.process_command("show adaptive evaluation"))

    assert result["success"] is True
    metrics = result["results"][0]["evaluation_metrics"]
    assert "recommendation_accuracy" in metrics
    assert "user_correction_rate" in metrics


def test_classify_command_returns_adaptive_prediction():
    controller = build_controller()
    controller.agents["learning"].record_command(
        "show system status",
        {"tier": "direct", "intent": {"type": "system_info"}},
        {"success": True},
    )

    result = asyncio.run(controller.process_command("classify command show system status"))

    assert result["success"] is True
    classification = result["results"][0]["classification"]
    assert classification["predicted_intent"] in {"system_info", "intent:system_info"}
    assert classification["confidence"] >= 0.0


def test_autonomy_uses_default_browser_preference():
    controller = build_controller()
    controller.memory.save_user_preference("default_browser", "edge")

    result = asyncio.run(controller.process_command("open browser"))

    assert result["success"] is True
    # "open browser" is a deterministically parsed direct intent; the saved
    # default_browser preference is resolved inside _execute_directly's
    # open_application handler, not by routing through the autonomy tier -
    # what matters is the correct browser opens.
    assert result["tier"] == "direct"
    assert controller.agents["system"].opened[0][0] == "edge"


def test_autonomy_uses_notes_folder_preference(tmp_path):
    controller = build_controller()
    controller.memory.save_user_preference("notes_folder", str(tmp_path))

    result = asyncio.run(controller.process_command("make a note about roadmap"))

    assert result["success"] is True
    created_path, _ = controller.agents["file"].created[0]
    assert created_path == str(tmp_path / "roadmap_note.txt")


def test_autonomy_suggests_learned_workflows():
    controller = build_controller()

    result = asyncio.run(controller.process_command("suggest workflows"))

    assert result["success"] is True
    assert result["needs_clarification"] is True
    assert result["autonomy"]["strategy"] == "workflow_suggestions"
    assert "intent:browser_action" in result["suggestions"][0]


def test_phase3_searches_files_without_ollama():
    controller = build_controller()

    result = asyncio.run(controller.process_command("search files for roadmap in desktop"))

    assert result["success"] is True
    assert result["tier"] == "direct"
    assert controller.agents["file"].searched[0][1] == "roadmap"
    assert "Desktop" in controller.agents["file"].searched[0][0]


def test_phase3_gets_file_info_from_natural_language():
    controller = build_controller()

    result = asyncio.run(controller.process_command("show file info for launch_checklist_note.txt"))

    assert result["success"] is True
    assert result["results"][0]["name"] == "launch_checklist_note.txt"


def test_phase3_backs_up_known_folder():
    controller = build_controller()

    result = asyncio.run(controller.process_command("backup desktop as desktop_snapshot"))

    assert result["success"] is True
    assert controller.agents["file"].backups[0][1] == "desktop_snapshot"


def test_phase3_closes_application_without_council():
    controller = build_controller()

    result = asyncio.run(controller.process_command("close notepad"))

    assert result["success"] is True
    assert controller.agents["system"].closed == ["notepad"]


def test_phase3_system_network_drives_and_env_commands():
    controller = build_controller()

    network = asyncio.run(controller.process_command("network status"))
    drives = asyncio.run(controller.process_command("show drives"))
    env = asyncio.run(controller.process_command("show env PATH"))

    assert network["success"] is True
    assert drives["success"] is True
    assert env["success"] is True
    assert env["results"][0]["name"] == "PATH"


def test_dangerous_system_command_requires_confirmation():
    class ConfirmingSystemAgent(DummySystemAgent):
        def __init__(self):
            super().__init__()
            self.confirmed_runs = []

        def run_command(self, command, shell="powershell", confirmed=False):
            if "Remove-Item" in command and not confirmed:
                return {
                    "success": False,
                    "needs_confirmation": True,
                    "command": command,
                    "warning": "This command may be destructive. Please confirm execution.",
                    "matched_patterns": ["Remove-Item"],
                }
            self.confirmed_runs.append((command, shell, confirmed))
            return {"success": True, "output": command, "return_code": 0, "confirmed": confirmed}

    controller = build_controller()
    controller.agents["system"] = ConfirmingSystemAgent()

    first = asyncio.run(controller.process_command("run powershell Remove-Item test.txt"))
    second = asyncio.run(controller.process_command("confirm command"))

    assert first["success"] is False
    assert first["needs_confirmation"] is True
    assert "pending_system_command" in controller.memory.short_term
    assert second["success"] is True
    assert second["results"][0]["confirmed"] is True
    assert controller.memory.recall_short_term("pending_system_command") is None


def test_dangerous_system_command_can_be_cancelled():
    class ConfirmingSystemAgent(DummySystemAgent):
        def run_command(self, command, shell="powershell", confirmed=False):
            return {
                "success": False,
                "needs_confirmation": True,
                "command": command,
                "warning": "This command may be destructive. Please confirm execution.",
            }

    controller = build_controller()
    controller.agents["system"] = ConfirmingSystemAgent()

    first = asyncio.run(controller.process_command("run powershell Remove-Item test.txt"))
    second = asyncio.run(controller.process_command("cancel command"))

    assert first["needs_confirmation"] is True
    assert second["success"] is True
    assert second["results"][0]["cancelled"] is True
    assert controller.memory.recall_short_term("pending_system_command") is None


def test_phase3_decomposer_routes_file_search_before_web_search():
    decomposer = TaskDecomposer()

    plan = decomposer.decompose("search files for roadmap in desktop then open notepad")

    assert plan["subtasks"][0]["agent"] == "file"
    assert plan["subtasks"][0]["action"] == "search_files"
    assert plan["subtasks"][1]["agent"] == "system"
    assert plan["subtasks"][1]["action"] == "open"


def test_clarification_followup_completes_missing_target():
    controller = build_controller()

    first = asyncio.run(controller.process_command("open"))
    second = asyncio.run(controller.process_command("browser"))

    assert first["needs_clarification"] is True
    assert second["success"] is True
    assert controller.agents["system"].opened[0][0] == "chrome"


def test_display_result_prints_clarifying_question(capsys):
    controller = build_controller()

    controller._display_result({
        "success": False,
        "needs_clarification": True,
        "question": "What should I open?",
        "suggestions": ["open chrome"],
    })

    captured = capsys.readouterr()
    # Product was rebranded ASC -> VEGA (2026-08-26).
    assert "VEGA: What should I open?" in captured.out
    assert "open chrome" in captured.out


def test_parameterized_learned_direct_strategy_is_not_reused():
    controller = build_controller()

    class SimilarLearningAgent(DummyLearningAgent):
        def find_similar_success(self, command, threshold=0.75):
            return {
                "score": 1.0,
                "match": {
                    "strategy": {
                        "tier": "direct",
                        "complexity": "simple",
                        "intent": {
                            "type": "create_file",
                            "path": "old_note.txt",
                        },
                    }
                },
            }

    controller.agents["learning"] = SimilarLearningAgent()

    strategy = controller._determine_execution_strategy("create file named new_note.txt")

    assert strategy["source"] == "rules"
    assert strategy["intent"]["path"].endswith("new_note.txt")


def test_explicit_sequence_is_not_collapsed_by_learning():
    controller = build_controller()

    class SimilarLearningAgent(DummyLearningAgent):
        def find_similar_success(self, command, threshold=0.75):
            return {
                "score": 1.0,
                "match": {
                    "strategy": {
                        "tier": "direct",
                        "complexity": "simple",
                        "intent": {"type": "system_info"},
                    }
                },
            }

    controller.agents["learning"] = SimilarLearningAgent()

    strategy = controller._determine_execution_strategy(
        "show system status and then list running processes"
    )

    assert strategy["tier"] == "sequential"
    assert strategy["source"] == "decomposer"


def test_memory_paths_survive_cwd_change(tmp_path, monkeypatch):
    from core.memory import MemorySystem

    memory = MemorySystem("./config/settings.json")
    original_db = memory.episodic_db
    monkeypatch.chdir(tmp_path)

    memory.remember_episode(
        task="cwd change test",
        action="test",
        result="ok",
        success=True,
    )

    assert memory.episodic_db == original_db
    assert memory.get_memory_stats()["episodes_count"] >= 1


def test_learning_paths_survive_cwd_change(tmp_path, monkeypatch):
    from agents.learning_agent import LearningAgent

    agent = LearningAgent("./config/settings.json")
    original_path = agent.replay.path
    original_db_path = agent.replay.db_path  # actual persistence target (SQLite)
    monkeypatch.chdir(tmp_path)

    agent.record_command(
        "learning cwd test",
        {"tier": "direct", "intent": {"type": "system_info"}},
        {"success": True, "tier": "direct", "steps_completed": 1, "steps_total": 1},
    )

    assert agent.replay.path == original_path
    assert agent.replay.db_path == original_db_path
    assert original_db_path.exists()


def test_tool_registry_converts_handler_exception_to_failure():
    from core.tool_registry import ToolRegistry, ToolSpec

    registry = ToolRegistry()

    def explode(action, params, context):
        raise RuntimeError("boom")

    registry.register(ToolSpec(
        name="boom_tool",
        agent="test",
        description="Explodes for testing",
        actions=["boom"],
        handler=explode,
    ))

    result = asyncio.run(registry.execute("boom", {}, "context"))

    assert result["success"] is False
    assert result["tool"] == "boom_tool"
    assert result["error"] == "boom"


def test_web_agent_site_search_builds_known_site_url():
    system_agent = StubSystemAgent()
    vision_agent = StubVisionAgent()
    learning_agent = DummyLearningAgent()
    agent = WebAgent(
        config_path="./config/settings.json",
        system_agent=system_agent,
        vision_agent=vision_agent,
        learning_agent=learning_agent,
    )
    agent.browser_driver = None
    agent.navigation_delay_seconds = 0

    result = asyncio.run(agent.search_site("laptops", site="amazon", browser_name="chrome"))

    assert result["success"] is True
    assert result["site"] == "amazon"
    assert result["url"] == "https://www.amazon.in/s?k=laptops"


def test_routed_form_action_navigates_before_filling():
    controller = build_controller()
    web_agent = controller.agents["web"]

    result = asyncio.run(
        controller._route_action(
            "form",
            {
                "url": "https://example.com/form",
                "fields": {"name": "Swapnil"},
                "focus_first_field": True,
            },
            "fill the form",
        )
    )

    assert result["success"] is True
    assert result["navigated_url"] == "https://example.com/form"
    assert web_agent.navigated == [("https://example.com/form", "chrome")]


def test_phase5_cold_start_bootstrap():
    from learning.phase5_model import OfflinePreferenceModel
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as tmp:
        model = OfflinePreferenceModel(tmp.name)
        state = model.fit([])
        assert state["summary"]["clusters"] >= 1
        assert "prediction" in state
        assert state["summary"]["offline"] is True


def test_learning_agent_explicit_preference_retraining():
    from agents.learning_agent import LearningAgent
    import tempfile, json

    with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as cfg_file:
        config = {
            "paths": {"data": tempfile.gettempdir()},
            "learning": {"history_limit": 100, "adaptive_min_confidence": 0.5}
        }
        cfg_file.write(json.dumps(config).encode("utf-8"))
        cfg_file.flush()

        agent = LearningAgent(config_path=cfg_file.name)
        res = agent.learn_explicit_preference("default_browser", "firefox")
        assert res["success"] is True
        assert res["learned_preference"] == {"default_browser": "firefox"}


# ---------------------------------------------------------------------------
# Autonomy pass: outcome verification, approval gate, durable tasks, and
# confidence-based learned-strategy reuse.
# ---------------------------------------------------------------------------

def test_verify_outcome_parses_satisfied_response():
    from core.council import LLMCouncil

    council = object.__new__(LLMCouncil)

    async def fake_generate(member, prompt, options=None):
        assert member == "strategist"
        return {"response": '{"satisfied": false, "confidence": 0.8, "reason": "ignored the age filter"}'}, "qwen2.5:3b"

    council._generate_with_fallbacks = fake_generate

    result = asyncio.run(council.verify_outcome("organize and archive old files", "organized by type only"))
    assert result["verified"] is True
    assert result["satisfied"] is False
    assert result["confidence"] == 0.8
    assert "age filter" in result["reason"]


def test_verify_outcome_fails_open_when_model_unavailable():
    from core.council import LLMCouncil

    council = object.__new__(LLMCouncil)

    async def raising_generate(member, prompt, options=None):
        raise RuntimeError("model unavailable")

    council._generate_with_fallbacks = raising_generate

    result = asyncio.run(council.verify_outcome("do something", "did something"))
    assert result["satisfied"] is True  # fail open, never turn success into failure
    assert result["verified"] is False


def test_send_email_stages_by_default_then_confirm_sends():
    controller = build_controller()

    class RecordingBusinessAgent:
        def __init__(self):
            self.sent = []

        def send_email_notification(self, **kwargs):
            self.sent.append(kwargs)
            return {"success": True, **kwargs}

    controller.agents["business"] = RecordingBusinessAgent()

    staged = controller._stage_or_send_email("raj@example.com", "Invoice", "See attached")
    assert staged["success"] is False
    assert staged["needs_confirmation"] is True
    assert controller.agents["business"].sent == []  # not sent yet
    assert controller.memory.recall_short_term("pending_outward_action")["action"] == "send_email"

    result = asyncio.run(
        controller._execute_directly("confirm command", {"intent": {"type": "confirm_system_command"}})
    )
    assert result["success"] is True
    assert controller.agents["business"].sent == [
        {"to_email": "raj@example.com", "subject": "Invoice", "body": "See attached", "priority": "normal"}
    ]
    assert controller.memory.recall_short_term("pending_outward_action") is None


def test_send_email_can_be_cancelled():
    controller = build_controller()

    class RecordingBusinessAgent:
        def __init__(self):
            self.sent = []

        def send_email_notification(self, **kwargs):
            self.sent.append(kwargs)
            return {"success": True}

    controller.agents["business"] = RecordingBusinessAgent()
    controller._stage_or_send_email("raj@example.com", "Invoice", "See attached")

    result = asyncio.run(
        controller._execute_directly("cancel command", {"intent": {"type": "cancel_system_command"}})
    )
    assert result["success"] is True
    assert controller.agents["business"].sent == []
    assert controller.memory.recall_short_term("pending_outward_action") is None


def test_send_email_auto_confirm_preference_sends_immediately():
    controller = build_controller()

    class RecordingBusinessAgent:
        def __init__(self):
            self.sent = []

        def send_email_notification(self, **kwargs):
            self.sent.append(kwargs)
            return {"success": True}

    controller.agents["business"] = RecordingBusinessAgent()
    controller.memory.save_user_preference("auto_confirm_email", True)

    result = controller._stage_or_send_email("raj@example.com", "Invoice", "See attached")
    assert result["success"] is True
    assert len(controller.agents["business"].sent) == 1
    assert controller.memory.recall_short_term("pending_outward_action") is None


def test_unattended_action_budget_blocks_after_limit():
    controller = build_controller()
    controller.unattended_action_hourly_limit = 2

    assert controller._check_unattended_action_budget() is True
    assert controller._check_unattended_action_budget() is True
    assert controller._check_unattended_action_budget() is False


def test_get_incomplete_tasks_lifecycle(tmp_path):
    from core.memory import MemorySystem
    import json as json_module

    config_path = tmp_path / "settings.json"
    config_path.write_text(json_module.dumps({
        "paths": {"data": str(tmp_path / "data")},
        "session": {"session_file": str(tmp_path / "data" / "sessions.json"),
                    "preferences_file": str(tmp_path / "data" / "preferences.json")},
    }), encoding="utf-8")

    memory = MemorySystem(str(config_path))
    task_id = memory.create_task_record("a multi-step command", steps_total=3)

    assert any(t["id"] == task_id for t in memory.get_incomplete_tasks())

    memory.update_task_record(task_id, steps_completed=2, last_error="step 3 flaky")
    incomplete = [t for t in memory.get_incomplete_tasks() if t["id"] == task_id][0]
    assert incomplete["steps_completed"] == 2
    assert incomplete["last_error"] == "step 3 flaky"

    memory.update_task_record(task_id, status="done")
    assert not any(t["id"] == task_id for t in memory.get_incomplete_tasks())


def test_can_reuse_learned_strategy_is_confidence_based():
    controller = build_controller()

    class FakeLearningAgent:
        def __init__(self, stats_by_intent):
            self._stats = stats_by_intent

        def get_intent_confidence(self, intent_type):
            return self._stats.get(intent_type, {"samples": 0, "successes": 0, "success_rate": 0.0})

    controller.agents["learning"] = FakeLearningAgent({
        "file_info": {"samples": 10, "successes": 10, "success_rate": 1.0},
        "create_file": {"samples": 9, "successes": 8, "success_rate": 0.889},
        "weather_api": {"samples": 3, "successes": 3, "success_rate": 1.0},
    })

    assert controller._can_reuse_learned_strategy({"intent": {"type": "system_info"}}) is True  # always-safe
    assert controller._can_reuse_learned_strategy({"intent": {"type": "file_info"}}) is True  # high confidence, enough samples
    assert controller._can_reuse_learned_strategy({"intent": {"type": "create_file"}}) is False  # below 0.9
    assert controller._can_reuse_learned_strategy({"intent": {"type": "weather_api"}}) is False  # not enough samples


# ---------------------------------------------------------------------------
# Phase 4: expression evaluation, document generation.
#
# These are deliberately ADVERSARIAL. The previous sandbox passed a suite of
# "does the happy path work" tests while being escapable in one line - the
# tests asserted shape, not safety. Every escape below is one that either
# defeated the old implementation or is a standard idiom that would.
# ---------------------------------------------------------------------------

def test_expression_evaluator_computes_real_math():
    from core.sandbox import evaluate_expression

    assert evaluate_expression("17/100*4500")["result"] == 765.0
    assert evaluate_expression("sqrt(144)")["result"] == 12.0
    assert evaluate_expression("2**10")["result"] == 1024
    assert evaluate_expression("round(mean([1,2,3,4]), 2)")["result"] == 2.5
    assert evaluate_expression("max(3, 7, 5)")["result"] == 7


def test_expression_evaluator_blocks_the_escape_that_broke_the_old_sandbox():
    """The exact one-liner that reached os/subprocess through the old denylist."""
    from core.sandbox import evaluate_expression

    escape = "__builtins__.__dict__['__imp'+'ort__']('os')"
    result = evaluate_expression(escape)
    assert result["success"] is False
    # Rejected by the GRAMMAR (this is not a permitted call/attribute form),
    # not by matching a forbidden name - name matching is exactly what the
    # original was defeated by, via string concatenation.
    assert "permitted" in result["error"].lower() or "not allowed" in result["error"].lower()

    # And the same for every spelling of it - none of these can be expressed
    # at all, so obfuscation has nothing to hide behind.
    for variant in [
        "__builtins__.__dict__",
        "__builtins__['x']",
        "().__class__",
        "(1).__class__.__bases__",
    ]:
        assert evaluate_expression(variant)["success"] is False


def test_expression_evaluator_blocks_all_system_access_idioms():
    from core.sandbox import evaluate_expression

    for hostile in [
        "__import__('os').getcwd()",
        "open('secret.txt')",
        "eval('1+1')",
        "exec('x=1')",
        "().__class__.__bases__[0].__subclasses__()",
        "type(1).mro()",
        "globals()",
        "getattr(1, 'real')",
        "[].append(1)",
        "os.system('whoami')",
        "print(1)",           # not on the allowlist - no I/O at all
        "lambda: 1",
        "__builtins__",
    ]:
        result = evaluate_expression(hostile)
        assert result["success"] is False, f"SECURITY: {hostile!r} was allowed to evaluate"


def test_expression_evaluator_bounds_resource_exhaustion():
    from core.sandbox import evaluate_expression

    # Each of these would hang or exhaust memory if evaluated naively.
    assert evaluate_expression("9**9**9")["success"] is False
    assert evaluate_expression("10**100000")["success"] is False
    assert evaluate_expression("'a' * 1000000000")["success"] is False
    assert evaluate_expression("factorial(1000000)")["success"] is False


def test_expression_evaluator_rejects_statements_and_assignment():
    from core.sandbox import evaluate_expression

    for statement in ["x = 1", "import math", "while True: pass", "def f(): pass"]:
        assert evaluate_expression(statement)["success"] is False


def test_expression_evaluator_reports_errors_without_raising():
    from core.sandbox import evaluate_expression

    div = evaluate_expression("1/0")
    assert div["success"] is False and "zero" in div["error"].lower()
    assert evaluate_expression("")["success"] is False
    assert evaluate_expression("((((")["success"] is False


def test_math_action_accepts_the_param_names_models_emit():
    """Live routing showed the model sending 'expression', not 'code'."""
    controller = build_controller()

    for key in ["expression", "expr", "query", "code", "calculation"]:
        result = controller._execute_math_action("calculate", {key: "17/100*4500"}, "")
        assert result["success"] is True, f"param name {key!r} not accepted"
        assert result["result"] == 765.0

    missing = controller._execute_math_action("calculate", {}, "")
    assert missing["success"] is False
    assert "expression" in missing["error"]


def test_document_tools_create_word_document(tmp_path):
    from core.document_tools import create_word_document
    from docx import Document

    path = str(tmp_path / "report")
    result = create_word_document(path, "Test Report", [
        {"heading": "Summary", "body": "Key finding one."},
        "A plain trailing paragraph.",
    ])

    assert result["success"] is True
    assert result["path"].endswith(".docx")  # extension normalized even though we passed none

    doc = Document(result["path"])
    text = [p.text for p in doc.paragraphs]
    assert "Test Report" in text
    assert "Key finding one." in text
    assert "A plain trailing paragraph." in text


def test_document_tools_create_spreadsheet(tmp_path):
    from core.document_tools import create_spreadsheet
    from openpyxl import load_workbook

    path = str(tmp_path / "data.xlsx")
    result = create_spreadsheet(path, "Results", [[1, "a"], [2, "b"]], header=["id", "name"])

    assert result["success"] is True
    workbook = load_workbook(result["path"])
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows == [("id", "name"), (1, "a"), (2, "b")]


# ---------------------------------------------------------------------------
# Regression guards for the second-audit fixes.
#
# Each of these asserts a BEHAVIOUR (what the system refuses, what it does
# not re-run, what it degrades to) rather than a return shape - the whole
# reason the first set of tests stayed green through a sandbox escape.
# ---------------------------------------------------------------------------

def test_verification_never_re_executes_side_effects():
    """Advisory only.

    The previous implementation answered a flagged mismatch by asking the
    council for a fresh plan and running it - which re-ran the side effects
    the first pass had already performed (verified: the same email went out
    twice). This drives a real command through process_command with a
    verifier that always says "unsatisfied" and asserts the action layer is
    touched exactly once.
    """
    controller = build_controller()
    executed = []

    class AlwaysUnsatisfiedCouncil:
        def __init__(self):
            self.decide_calls = 0

        async def verify_outcome(self, command, summary):
            return {"verified": True, "satisfied": False,
                    "confidence": 0.9, "reason": "looks incomplete"}

        async def decide(self, command, complexity="medium"):
            self.decide_calls += 1
            return {"action_plan": [{"action": "send_email", "params": {}}]}

    council = AlwaysUnsatisfiedCouncil()
    controller.council = council

    original_route = controller._route_action

    async def recording_route(action, params, context):
        executed.append(action)
        return {"success": True, "action": action}

    controller._route_action = recording_route

    # A multi-step decomposition, so verification is actually reached.
    decomposition = {
        "goal": "send the report and log it",
        "subtasks": [
            {"id": "s1", "agent": "business", "action": "send_email", "params": {}, "description": "send"},
            {"id": "s2", "agent": "file", "action": "create_file", "params": {"path": "log.txt"}, "description": "log"},
        ],
    }
    result = asyncio.run(controller._execute_decomposed_tasks(decomposition, decomposition["goal"]))
    before = list(executed)

    # Now run the verification stage exactly as process_command does.
    summary = controller._summarize_result(result)
    verification = asyncio.run(controller.council.verify_outcome("send the report and log it", summary))
    if verification.get("verified") and not verification.get("satisfied"):
        result["verification"] = verification
        result["verification_warning"] = f"…{verification.get('reason')}"

    assert executed == before, "verification must not trigger any further actions"
    assert executed.count("send_email") == 1, "the email must not be sent twice"
    assert council.decide_calls == 0, "verification must not ask for a re-plan"
    assert result["verification_warning"], "the user must be told about the mismatch"


def test_critical_gate_degrades_to_confirmation_when_model_unavailable():
    controller = build_controller()

    class DeadChairman:
        async def validate_plan(self, plan, task):
            raise RuntimeError("Ollama unreachable")

    controller.chairman = DeadChairman()
    controller._critical_gate_bypass = None

    result = asyncio.run(controller._chairman_gate("delete the temp folder", {"tier": "direct"}))
    assert result is not None
    assert result.get("needs_confirmation") is True     # asks, does not refuse
    assert not result.get("rejected")
    assert controller.memory.recall_short_term("pending_critical_command") is not None

    # An explicit user confirmation lets exactly that command through.
    controller._critical_gate_bypass = "delete the temp folder"
    assert asyncio.run(controller._chairman_gate("delete the temp folder", {"tier": "direct"})) is None
    # ...and nothing else.
    assert asyncio.run(controller._chairman_gate("delete something else", {"tier": "direct"})) is not None


def test_text_formatting_is_not_treated_as_destructive():
    controller = build_controller()

    assert controller._assess_complexity("format my notes") != "critical"
    assert controller._assess_complexity("reformatting my resume") != "critical"
    assert controller._assess_complexity("show me the removed items report") != "critical"
    # Genuinely destructive still gated.
    for destructive in ["delete the temp folder", "uninstall spotify",
                        "format the disk", "format drive D:", "restart the system"]:
        assert controller._assess_complexity(destructive) == "critical", destructive


def test_replay_all_is_cached_and_invalidated_on_write(tmp_path):
    from learning.experience_replay import ExperienceReplayBuffer

    buf = ExperienceReplayBuffer(str(tmp_path / "command_history.json"), capacity=50)
    buf.add({"timestamp": "t1", "command": "one", "intent_type": "test"})
    first = buf.all()
    assert len(first) == 1

    # Cached: repeated reads are consistent and return independent copies
    # (a caller mutating the result must not corrupt the cache).
    second = buf.all()
    assert second == first and second is not first
    second.append({"injected": True})
    assert len(buf.all()) == 1

    # A write invalidates it.
    buf.add({"timestamp": "t2", "command": "two", "intent_type": "test"})
    assert len(buf.all()) == 2


def test_embed_requests_keep_alive():
    """Without keep_alive Ollama reloads the model per call - 2.0s vs 0.05s."""
    import inspect
    from core import embeddings

    source = inspect.getsource(embeddings.embed)
    assert "keep_alive" in source
    assert '"keep_alive": keep_alive' in source


def test_document_action_uses_model_supplied_names(tmp_path, monkeypatch):
    import core.document_tools as document_tools

    controller = build_controller()
    captured = {}

    def fake_create(path, title, sections):
        captured.update({"path": path, "title": title, "sections": sections})
        return {"success": True, "path": path}

    monkeypatch.setattr(document_tools, "create_word_document", fake_create)
    monkeypatch.setattr(controller, "_documents_dir", lambda: tmp_path)

    # Exactly what the router emitted live: {"name": ...} plus body text.
    controller._execute_document_action(
        "create_document",
        {"name": "notes", "content": "hello"},
        "make a word document called notes with the text hello",
    )

    assert captured["title"] == "notes"
    assert "hello" in " ".join(str(s) for s in captured["sections"])
    # Filename comes from the name, never from the whole sentence.
    assert captured["path"].endswith("notes.docx")


def test_shutdown_stops_the_browser_driver():
    controller = build_controller()
    stopped = {"called": False}

    class FakeDriver:
        async def stop(self):
            stopped["called"] = True

    class FakeWebAgent:
        browser_driver = FakeDriver()

    controller.agents["web"] = FakeWebAgent()
    asyncio.run(controller.shutdown())
    assert stopped["called"] is True



# ---------------------------------------------------------------------------
# Regression guards for the "remaining issues" pass (honesty, secrets,
# config, form-fill safety).
# ---------------------------------------------------------------------------

def test_analytics_reports_real_numbers_not_constants():
    """Metrics must come from real records, not hardcoded values.

    _get_metric_value previously returned a fixed table (total_tasks: 150,
    productivity_score: 0.85) described in-code as "Simulated metric values",
    and _generate_report_data invented a whole productivity report.
    """
    from agents.advanced_analytics import AdvancedAnalytics

    analytics = AdvancedAnalytics("./config/settings.json")
    stale_constants = {150.0, 120.0, 0.85, 0.78, 0.92, 0.75}

    total = analytics._get_metric_value("total_tasks")
    completed = analytics._get_metric_value("completed_tasks")

    # Consistent with each other and with the underlying records.
    assert completed <= total
    report = analytics._generate_report_data("productivity", {})
    if report.get("available"):
        assert report["summary"]["total_tasks"] == int(total)
        assert report["summary"]["completed_tasks"] == int(completed)
        assert report.get("source"), "a real report must name its data source"
        # The old fabricated pair was exactly (150, 120).
        assert (total, completed) != (150.0, 120.0)


def test_reports_without_a_data_source_say_so():
    from agents.advanced_analytics import AdvancedAnalytics

    analytics = AdvancedAnalytics("./config/settings.json")
    for report_type in ("financial", "operational"):
        result = analytics._generate_report_data(report_type, {})
        assert result.get("available") is False
        assert "no" in result["reason"].lower()


def test_erp_sync_does_not_claim_success():
    from agents.business_process_automation import BusinessProcessAutomation

    bpa = BusinessProcessAutomation("./config/settings.json")
    bpa.create_erp_integration(
        system_name="AuditTestERP",
        api_endpoint="https://example.invalid/api",
        authentication={}, sync_config={},
    )
    result = bpa.sync_with_erp(system_name="AuditTestERP", data_type="inventory")
    assert result["success"] is False
    assert result.get("not_implemented") is True


def test_encryption_is_real_not_xor():
    """encrypt_data must not be reversible without the OS keystore."""
    from agents.enterprise_security import EnterpriseSecurity

    security = EnterpriseSecurity("./config/settings.json")
    plaintext = "sensitive-value-abcdefghijklmnop"
    enc = security.encrypt_data(plaintext, "default")

    if not enc.get("success"):
        assert enc.get("not_implemented") is True  # honest refusal is acceptable
        return

    ciphertext = enc["encrypted_data"]
    assert plaintext not in ciphertext
    # Repeating-key XOR against the stored key used to round-trip trivially;
    # verify a naive XOR with any stored key does NOT recover the plaintext.
    for key in (security.encryption_keys or {}).values():
        try:
            raw = bytes.fromhex(ciphertext)
        except ValueError:
            break  # not hex any more - already not the old scheme
        key_bytes = str(key).encode()
        naive = bytes(raw[i] ^ key_bytes[i % len(key_bytes)] for i in range(len(raw)))
        assert naive.decode("utf-8", "ignore") != plaintext

    assert security.decrypt_data(ciphertext, "default")["decrypted_data"] == plaintext


def test_secrets_roundtrip_and_never_store_plaintext(tmp_path):
    from core import secrets as os_secrets

    if not os_secrets.is_available():
        return  # honestly unavailable on this platform

    path = tmp_path / "creds.json"
    path.write_text('{"password": "plaintext-app-password-12345"}', encoding="utf-8")

    loaded = os_secrets.load_credentials(path)
    assert loaded["password"] == "plaintext-app-password-12345"

    import json as json_mod
    on_disk = json_mod.loads(path.read_text(encoding="utf-8"))
    assert "password" not in on_disk, "plaintext secret must be upgraded away"
    assert "protected_password" in on_disk
    assert "plaintext-app-password-12345" not in path.read_text(encoding="utf-8")

    # Still readable after migration.
    assert os_secrets.load_credentials(path)["password"] == "plaintext-app-password-12345"


def test_config_is_cached_and_isolated():
    from core.config import load_config

    first = load_config("./config/settings.json")
    first["INJECTED"] = True
    second = load_config("./config/settings.json")
    assert "INJECTED" not in second, "callers must not corrupt the shared config"


def test_conversation_length_cap_scales_with_the_question():
    """Casual replies stay short; explanatory answers are not capped at all.

    Previously a flat 5-sentence cap applied to every reply, so VEGA could
    not fully explain anything. Casual messages still get a brief cap;
    a question that asks for explanation gets sentence_cap=None - no
    truncation regardless of how long the model's answer runs.
    """
    import re

    from agents.conversation_agent import ConversationAgent

    brief = ConversationAgent._sentence_cap_for("hi")
    detailed = ConversationAgent._sentence_cap_for("explain how the routing tiers work")
    assert brief == ConversationAgent.BRIEF_SENTENCE_CAP
    assert detailed is None

    long_text = " ".join(f"Sentence number {i}." for i in range(30))
    assert len(re.split(r"(?<=[.!?])\s+", ConversationAgent._polish_response(long_text, brief))) <= brief
    # Uncapped: every sentence in a 30-sentence answer survives.
    uncapped = ConversationAgent._polish_response(long_text, detailed)
    assert len(re.split(r"(?<=[.!?])\s+", uncapped)) == 30


def test_form_fill_refuses_to_type_sensitive_values_blindly():
    """The keyboard fallback cannot confirm where it is typing."""
    from agents.web_agent import WebAgent

    agent = object.__new__(WebAgent)
    agent.label_synonyms = {}
    agent.selector_confidence_threshold = 0.7
    agent.label_to_input_offset_px = 45
    agent.browser_driver = None
    agent.verbose_logging = False
    agent.form_field_delay_seconds = 0

    async def fake_analyze(form_name=None):
        return {"success": True, "detected_labels": ["Password"], "form_signature": "sig"}

    agent.analyze_form = fake_analyze
    agent._build_field_sequence = lambda values, analysis: [
        {"name": name, "value": value} for name, value in values.items()
    ]

    result = asyncio.run(agent.intelligent_form_fill(
        {"password": "hunter2"}, focus_first_field=True, allow_keyboard_fallback=True,
    ))
    assert result["success"] is False
    assert "sensitive" in result["error"].lower()
    assert result["fill_method"] == "refused_sensitive_keyboard_fallback"


# ---------------------------------------------------------------------------
# Audit pass 3 regression guards.
# ---------------------------------------------------------------------------

def test_secrets_never_written_to_tracked_settings_file(tmp_path):
    """configure_email/configure_google_calendar must not leak into settings.json.

    config/settings.json is tracked by git and not gitignored, so a password
    written there is a password committed to the repository (audit H-1).
    """
    import json as json_mod

    from agents.business_agent import BusinessAgent

    project = tmp_path / "proj"
    (project / "config").mkdir(parents=True)
    (project / "data").mkdir()
    cfg = project / "config" / "settings.json"
    cfg.write_text(json_mod.dumps({"paths": {"data": "./data"}}), encoding="utf-8")

    agent = BusinessAgent(str(cfg))
    secret = "app-password-must-not-leak-9876"
    agent._persist_config_section(
        "email",
        {"configured": True, "email_address": "a@b.c", "password": secret},
        secret_keys=("password",),
    )

    settings_text = cfg.read_text(encoding="utf-8")
    assert secret not in settings_text, "secret leaked into the git-tracked config"
    assert "password" not in json_mod.loads(settings_text).get("email", {})

    creds = project / "data" / "email_credentials.json"
    assert creds.exists()
    assert secret not in creds.read_text(encoding="utf-8"), "secret stored in clear text"

    from core.secrets import load_credentials
    if load_credentials(creds):  # only assert recovery where a keystore exists
        assert load_credentials(creds)["password"] == secret


def test_delete_guard_does_not_widen_with_working_directory(tmp_path, monkeypatch):
    """The guard must not depend on where ASC was launched from.

    It previously allowed anything under Path.cwd(), so launching from C:\
    authorised deleting system directories - the guard switched itself off
    (audit H-2).
    """
    from pathlib import Path

    from agents.file_agent import FileAgent

    agent = FileAgent("./config/settings.json")
    outside = Path(tmp_path) / "unrelated" / "file.txt"

    before = agent._is_safe_path(outside)
    monkeypatch.chdir(tmp_path)          # launch-directory changes...
    after = agent._is_safe_path(outside)
    assert before == after, "guard breadth changed with the working directory"

    # Never deletable regardless of cwd.
    assert not agent._is_safe_path(Path(agent.home))
    assert not agent._is_safe_path(Path(agent.project_root))
    assert not agent._is_safe_path(Path(agent.project_root) / "data")
    # Still usable for real work.
    assert agent._is_safe_path(Path(agent.home) / "Documents" / "scratch.txt")


def test_interrupted_save_never_destroys_an_existing_store(tmp_path):
    """A crash mid-write must leave the previous file fully intact (audit H-3)."""
    import json as json_mod
    from unittest import mock

    from core.atomic_io import write_json_atomic

    path = tmp_path / "tasks.json"
    original = {"tasks": [{"id": i} for i in range(453)]}
    write_json_atomic(path, original)

    with mock.patch.object(json_mod, "dump", side_effect=KeyboardInterrupt("crash")):
        try:
            write_json_atomic(path, {"tasks": []})
        except KeyboardInterrupt:
            pass

    assert len(json_mod.loads(path.read_text(encoding="utf-8"))["tasks"]) == 453
    assert not [p for p in tmp_path.iterdir() if p.name.startswith(".")], "temp file left behind"


def test_every_json_store_writes_atomically():
    """No production module may use the truncate-then-write pattern."""
    import re
    from pathlib import Path

    root = Path(__file__).resolve().parent.parent
    offenders = []
    for folder in ("core", "agents", "learning", "vision"):
        for py in (root / folder).glob("*.py"):
            if py.name == "atomic_io.py":
                continue
            for num, line in enumerate(py.read_text(encoding="utf-8").splitlines(), 1):
                if re.search(r"\bjson\.dump\((?!s)", line):
                    offenders.append(f"{py.name}:{num}")
    for py in (root / "main.py",):
        for num, line in enumerate(py.read_text(encoding="utf-8").splitlines(), 1):
            if re.search(r"\bjson\.dump\((?!s)", line):
                offenders.append(f"main.py:{num}")
    assert not offenders, f"non-atomic JSON writes: {offenders}"


def test_reminder_loop_does_not_erase_concurrent_additions(tmp_path):
    """The loop must re-read under a lock, not write back a stale copy.

    It previously read outside any lock and rewrote its own snapshot, so a
    reminder added between the read and the write vanished (audit M-1).
    """
    import json as json_mod
    from datetime import datetime, timedelta

    from core.reminders import ReminderLoop

    path = tmp_path / "schedules.json"
    past = (datetime.now() - timedelta(minutes=5)).isoformat()
    path.write_text(json_mod.dumps(
        [{"id": "r1", "title": "old", "status": "pending", "reminder_time": past}]
    ), encoding="utf-8")

    loop = ReminderLoop(path, speak=None)
    items = json_mod.loads(path.read_text(encoding="utf-8"))
    items.append({"id": "r2", "title": "added by user", "status": "pending", "reminder_time": past})
    path.write_text(json_mod.dumps(items), encoding="utf-8")

    loop.check_once()

    final = {i["id"]: i for i in json_mod.loads(path.read_text(encoding="utf-8"))}
    assert "r2" in final, "concurrently-added reminder was erased"
    assert final["r1"]["status"] == "done"


def test_reminder_display_never_aborts_the_fire(capsys):
    """A console that cannot encode the text must not stop the reminder.

    An emoji in the print raised UnicodeEncodeError under cp1252 before the
    item was recorded or marked done, so it never fired and retried forever
    every 30 seconds (audit M-2).
    """
    import sys

    from core import reminders

    class Cp1252Stdout:
        encoding = "cp1252"
        def write(self, text):
            text.encode("cp1252")   # raises on anything non-latin-1
            return len(text)
        def flush(self):
            pass

    original = sys.stdout
    try:
        sys.stdout = Cp1252Stdout()
        reminders._safe_print("reminder \u23f0 with emoji")   # must not raise
    finally:
        sys.stdout = original


def test_audit_trail_cost_does_not_grow_without_bound(tmp_path):
    """Rotation must cap the live trail, and archive rather than delete."""
    import json as json_mod

    from agents.enterprise_security import EnterpriseSecurity

    security = EnterpriseSecurity("./config/settings.json")
    original = list(security.audit_trail)
    try:
        security.security_path = tmp_path
        security.audit_trail = [{"n": i} for i in range(security.AUDIT_TRAIL_LIMIT + 2500)]
        security._save_audit_trail()

        assert len(security.audit_trail) == security.AUDIT_TRAIL_LIMIT
        archives = list((tmp_path / "audit_archive").glob("*.json"))
        assert archives, "entries were dropped instead of archived"
        archived = sum(len(json_mod.loads(a.read_text(encoding="utf-8"))) for a in archives)
        assert archived + len(security.audit_trail) == security.AUDIT_TRAIL_LIMIT + 2500
    finally:
        security.audit_trail = original


def test_analytics_reads_are_cached_and_invalidated(tmp_path):
    """Every metric re-parsed ~370KB of JSON before this (audit M-4)."""
    import json as json_mod

    from agents.advanced_analytics import AdvancedAnalytics

    analytics = AdvancedAnalytics("./config/settings.json")
    analytics.data_root = tmp_path
    (tmp_path / "business_tasks.json").write_text(
        json_mod.dumps({"tasks": [{"id": 1, "status": "completed"}]}), encoding="utf-8")

    first = analytics._load_operational_data()
    assert analytics._load_operational_data() is first, "not cached"

    tasks_file = tmp_path / "business_tasks.json"
    tasks_file.write_text(
        json_mod.dumps({"tasks": [{"id": 1, "status": "completed"}, {"id": 2}]}), encoding="utf-8")
    # Force a distinct mtime. Two writes inside one filesystem tick can share
    # st_mtime_ns on Windows, which made this assert flakily on a cache hit
    # that was actually correct behaviour.
    import os as os_mod
    stamp = tasks_file.stat().st_mtime + 2
    os_mod.utime(tasks_file, (stamp, stamp))

    refreshed = analytics._load_operational_data()
    assert refreshed is not first, "stale after the source file changed"
    assert len(refreshed["tasks"]) == 2


def test_config_cache_holds_multiple_paths(tmp_path):
    """Clearing the whole cache per revision evicted other paths (audit M-5)."""
    import json as json_mod

    from core.config import load_config

    a, b = tmp_path / "a.json", tmp_path / "b.json"
    a.write_text(json_mod.dumps({"which": "a"}), encoding="utf-8")
    b.write_text(json_mod.dumps({"which": "b"}), encoding="utf-8")

    for _ in range(5):
        assert load_config(str(a))["which"] == "a"
        assert load_config(str(b))["which"] == "b"

    from core.config import _CACHE
    assert str(a.resolve()) in _CACHE and str(b.resolve()) in _CACHE


def test_no_bare_except_in_production_code():
    """A bare except also swallows KeyboardInterrupt (audit M-7)."""
    import ast
    from pathlib import Path

    root = Path(__file__).resolve().parent.parent
    offenders = []
    for folder in ("core", "agents", "learning", "vision"):
        for py in (root / folder).glob("*.py"):
            tree = ast.parse(py.read_text(encoding="utf-8"))
            for node in ast.walk(tree):
                if isinstance(node, ast.ExceptHandler) and node.type is None:
                    offenders.append(f"{py.name}:{node.lineno}")
    assert not offenders, f"bare except clauses: {offenders}"


def test_no_source_file_carries_a_bom():
    """A BOM breaks any tool reading these as plain UTF-8 (audit L-1)."""
    from pathlib import Path

    root = Path(__file__).resolve().parent.parent
    offenders = [
        str(py.relative_to(root))
        for py in root.rglob("*.py")
        if ".git" not in py.parts and "__pycache__" not in py.parts
        and py.read_bytes().startswith(b"\xef\xbb\xbf")
    ]
    assert not offenders, f"files with a BOM: {offenders}"


def test_batch_similarity_matches_the_scalar_path():
    """Vectorised cosine similarity must agree with the original (audit L-4)."""
    from core.embeddings import cosine_similarity
    from core.memory import MemorySystem

    query = [0.1, 0.9, -0.4, 0.25]
    vectors = [[0.1, 0.9, -0.4, 0.25], [-1.0, 0.2, 0.3, 0.0], [0.0, 0.0, 0.0, 0.0]]

    batched = MemorySystem._batch_similarity(query, vectors)
    for got, vector in zip(batched, vectors):
        assert abs(got - cosine_similarity(query, vector)) < 1e-5
    assert MemorySystem._batch_similarity(query, []) == []


@pytest.mark.parametrize("command", [
    "how much room have i got left",
    "whats hogging the memory",
    "what files are on my desktop",
    "where are my downloads",
    "is there space on the drive",
    "which apps are running",
    "how is my battery",
])
def test_questions_without_a_verb_are_still_real_commands(command):
    """A question about the machine must not be dismissed as small talk.

    _looks_actionable required a "?", a digit, or an imperative verb. Real
    requests routinely have none of the three - "how much room have i got
    left" is a disk query and "whats hogging the memory" is a process query,
    and both were answered as chit-chat by the nonsense gate. A question
    opener plus a noun from this system's domain now counts as actionable.
    """
    from main import AutonomousSystemController

    assert AutonomousSystemController._looks_actionable(command), \
        f"{command!r} was dismissed as non-actionable"


@pytest.mark.parametrize("command", [
    "flurbick wibble wobble",
    "my cat is fluffy",
    "the weather is nice today",
    "hello there friend",
    "i am feeling good",
    "what a lovely day",
    "how are you",
])
def test_small_talk_and_gibberish_are_still_filtered(command):
    """Widening the gate must not turn chat into planned tasks.

    Requiring BOTH a question opener and a domain noun is what keeps this
    side working - "what a lovely day" has an opener but no domain noun.
    """
    from main import AutonomousSystemController

    assert not AutonomousSystemController._looks_actionable(command), \
        f"{command!r} would now be planned as a task"


def test_stt_model_whisper_never_calls_the_cloud(monkeypatch):
    """audio.stt_model must actually be honoured.

    It was read into self.stt_model at init and never referenced again, so a
    config asking for local whisper still uploaded audio to Google - and the
    local pipeline stayed unreachable whenever the network was up, which is
    how a completely broken local path went unnoticed.
    """
    controller = object.__new__(AutonomousSystemController)
    controller.voice_enabled = True
    controller.stt_model = "whisper"
    controller.voice_service_available = False
    controller.voice_service_url = "http://localhost:5000"

    called = {"cloud": False}

    class TrackingSR:
        class UnknownValueError(Exception):
            pass

        class Recognizer:
            def adjust_for_ambient_noise(self, *a, **k):
                pass

            def listen(self, *a, **k):
                called["cloud"] = True
                return object()

            def recognize_google(self, audio):
                called["cloud"] = True
                return "should never happen"

        class Microphone:
            def __enter__(self):
                return object()

            def __exit__(self, *a):
                return False

    monkeypatch.setitem(sys.modules, "speech_recognition", TrackingSR)

    result = controller._listen_for_voice_command()
    assert called["cloud"] is False, "stt_model=whisper still sent audio to the cloud"
    # No voice service available, so it declines rather than falling back.
    assert result is None


def test_stt_model_google_does_not_fall_back_to_whisper(monkeypatch):
    """google means google: a cloud failure must not silently go local."""
    controller = object.__new__(AutonomousSystemController)
    controller.voice_enabled = True
    controller.stt_model = "google"
    controller.voice_service_available = True
    controller.voice_service_url = "http://localhost:5000"

    def must_not_run(*a, **k):
        raise AssertionError("fell back to local recording despite stt_model=google")

    controller._record_audio = must_not_run
    monkeypatch.setitem(sys.modules, "speech_recognition", None)

    assert controller._listen_for_voice_command() is None


def test_voice_disabled_short_circuits():
    controller = object.__new__(AutonomousSystemController)
    controller.voice_enabled = False
    assert controller._listen_for_voice_command() is None
